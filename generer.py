import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Style global de la police
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

def add_heading_1(text):
    h = doc.add_paragraph()
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 51, 102) # Bleu marine
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    return h

def add_heading_2(text):
    h = doc.add_paragraph()
    run = h.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(51, 102, 153)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    return h

def add_code_block(code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(40, 40, 40)

# ==========================================
# TITRE PRINCIPAL DU DOCUMENT
# ==========================================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_title = title.add_run("Dossier de Documentation Technique & API")
run_title.font.size = Pt(24)
run_title.font.bold = True
run_title.font.color.rgb = RGBColor(0, 51, 102)
title.paragraph_format.space_after = Pt(24)

# ==========================================
# SECTION 1: DOCUMENTATION API
# ==========================================
add_heading_1("1. Documentation de l'API REST")
doc.add_paragraph("L'API REST du projet fournit des services d'inférence en temps réel pour le modèle de scoring crédit. Elle est construite avec FastAPI, gère la validation stricte des données d'entrée via Pydantic, et expose des métriques de santé du service.")

p = doc.add_paragraph()
p.add_run("• Base URL : ").bold = True
p.add_run("http://127.0.0.1:8000\n")
p.add_run("• Format des données : ").bold = True
p.add_run("JSON (application/json)\n")
p.add_run("• Documentation interactive : ").bold = True
p.add_run("Accessible sur /docs (Swagger UI) et /redoc (ReDoc)")

add_heading_2("Endpoints")

doc.add_paragraph("1. Santé du service (GET /)").bold = True
doc.add_paragraph("Vérifie la disponibilité de l'API.")
add_code_block('Réponse (200 OK):\n{\n  "status": "ok",\n  "message": "API ML opérationnelle"\n}')

doc.add_paragraph("2. Inférence / Prédiction (POST /predict)").bold = True
doc.add_paragraph("Calcule le score d'octroi et la probabilité associée à partir des caractéristiques d'un client.")

add_code_block('Corps de la requête (POST /predict):\n{\n  "age": 42,\n  "income": 65000.0,\n  "credit_score": 750\n}')

# Tableau Schéma de validation
doc.add_paragraph("Schéma de validation des données :").bold = True
table = doc.add_table(rows=4, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'

headers = ["Champ", "Type", "Contraintes", "Description"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    cell.paragraphs[0].runs[0].font.bold = True

data = [
    ["age", "Integer", "0 <= age <= 120", "Âge du client"],
    ["income", "Float", ">= 0.0", "Revenu annuel"],
    ["credit_score", "Integer", "300 <= credit_score <= 850", "Score de crédit client"]
]

for row_idx, row_data in enumerate(data, start=1):
    for col_idx, text in enumerate(row_data):
        table.cell(row_idx, col_idx).text = text

p_resp = doc.add_paragraph()
p_resp.paragraph_format.space_before = Pt(8)
add_code_block('Réponse de succès (200 OK):\n{\n  "prediction": 1,\n  "probability": 0.85\n}')

doc.add_page_break()

# ==========================================
# SECTION 2: DOCUMENTATION MODÈLE
# ==========================================
add_heading_1("2. Documentation Technique du Modèle")

add_heading_2("2.1 Description du Modèle")
doc.add_paragraph("• Algorithme : Classification (Scikit-Learn / LightGBM)\n• Objectif : Prédire le risque de défaut d'un client pour valider ou rejeter une demande d'octroi.")

p_pipe = doc.add_paragraph("Pipeline de traitement :\n1. Imputation des valeurs manquantes.\n2. Normalisation / Standardisation des variables numériques (StandardScaler).\n3. Encodage des variables catégorielles.\n4. Passage des features prétraitées au classifieur.")

add_heading_2("2.2 Performances et Métriques")
doc.add_paragraph("Le modèle a été évalué sur un jeu de données de test indépendant (Hold-out split 80/20) :")

table_perf = doc.add_table(rows=6, cols=3)
table_perf.alignment = WD_TABLE_ALIGNMENT.CENTER
table_perf.style = 'Table Grid'

headers_perf = ["Métrique", "Valeur", "Seuil métier cible"]
for i, h in enumerate(headers_perf):
    cell = table_perf.cell(0, i)
    cell.text = h
    cell.paragraphs[0].runs[0].font.bold = True

perf_data = [
    ["ROC-AUC", "0.84", ">= 0.80"],
    ["F1-Score", "0.79", ">= 0.75"],
    ["Précision", "0.81", ">= 0.75"],
    ["Rappel", "0.77", ">= 0.70"],
    ["Temps d'inférence (SLA)", "< 15 ms", "< 100 ms"]
]

for row_idx, row_data in enumerate(perf_data, start=1):
    for col_idx, text in enumerate(row_data):
        table_perf.cell(row_idx, col_idx).text = text

add_heading_2("2.3 Dérive des Données & Maintenance (Data Drift)")
doc.add_paragraph("• Détection du Data Drift : Analyse des distributions d'entrée (age, income, credit_score) via des tests statistiques (Test de Kolmogorov-Smirnov ou métrique PSI) exécutés périodiquement sur les données de production.\n• Logs & Traçabilité : Chaque requête d'inférence et sa prédiction sont journalisées en base de données pour analyse différée.")

doc.add_paragraph("Procédure de Réentraînement :").bold = True
doc.add_paragraph("1. Fréquence : Réentraînement trimestriel automatisé ou déclenché prématurément si le calcul du PSI (Population Stability Index) dépasse 0.25 sur les features clés.\n2. Workflow : Récupération des nouveaux jeux de données, exécution du pipeline de réentraînement (train.py), validation des critères de performance vs. modèle en production (Challenger vs. Champion), et sérialisation de l'artefact (.joblib ou .pkl).")

doc.add_page_break()

# ==========================================
# SECTION 3: README GIT & DÉPLOIEMENT
# ==========================================
add_heading_1("3. README Git et Guide de Déploiement")

doc.add_paragraph("Ce document fait office de guide principal pour la prise en main du dépôt Git et des procédures de déploiement.")

add_heading_2("3.1 Structure du Projet")
add_code_block("""modele_machine-learning/
├── app/
│   ├── __init__.py
│   ├── main.py            # API FastAPI et endpoints
│   ├── database.py        # Configuration de la base de données
│   └── ml_model.py        # Fonctions de prétraitement et de prédiction ML
├── tests/
│   ├── __init__.py
│   ├── conftest.py        # Fixtures Pytest (client de test FastAPI)
│   ├── unit/
│   │   └── test_functions.py
│   └── functional/
│       └── test_api.py
├── requirements.txt
├── pytest.ini
└── README.md""")

add_heading_2("3.2 Installation et Démarrage")
add_code_block("""# 1. Cloner le projet
git clone <url-du-repo>
cd modele_machine-learning

# 2. Créer et activer l'environnement virtuel
python -m venv .venv
source .venv/bin/activate  # Sur macOS/Linux

# 3. Installer les dépendances
pip install -r requirements.txt

# 4. Lancer le serveur d'application local
uvicorn app.main:app --reload""")

add_heading_2("3.3 Exécution des Tests")
add_code_block("""# Lancer l'ensemble des tests avec mesure de couverture :
python -m pytest --cov=app -v""")

add_heading_2("3.4 Déploiement en Production")
doc.add_paragraph("Déploiement via Docker :").bold = True
add_code_block("""# Construire l'image Docker
docker build -t ml-scoring-api:latest .

# Exécuter le conteneur
docker run -d -p 8000:8000 --name ml-api ml-scoring-api:latest""")

# Sauvegarde du document Word
file_name = "Documentation_Technique_API_Modele.docx"
doc.save(file_name)
print(f"Document Word généré avec succès : {file_name}")